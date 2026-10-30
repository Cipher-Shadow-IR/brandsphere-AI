"""
Tests for native multi-page binary document parser: PDF, DOCX, and PPTX.
"""

import io
import pytest
from fastapi.testclient import TestClient
from backend.app.main import app
from backend.app.document_parser import (
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_pptx,
    parse_document_file
)

client = TestClient(app)


def test_docx_parser():
    import docx
    doc = docx.Document()
    doc.add_heading("Cloud Strategy Whitepaper", level=1)
    doc.add_paragraph("BrandSphere is guaranteed 100% cloud platform with unbeatable price.")
    table = doc.add_table(rows=1, cols=2)
    row = table.rows[0]
    row.cells[0].text = "Metric A"
    row.cells[1].text = "Value #1 in industry"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    parsed = extract_text_from_docx(file_bytes)
    assert parsed["format"] == "docx"
    assert "BrandSphere is guaranteed 100%" in parsed["text"]
    assert "Metric A | Value #1 in industry" in parsed["text"]
    assert parsed["table_count"] == 1


def test_pptx_parser():
    import pptx
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Executive Pitch Deck"
    slide.placeholders[1].text = "Our revolutionary engine never makes mistakes."

    buf = io.BytesIO()
    prs.save(buf)
    file_bytes = buf.getvalue()

    parsed = extract_text_from_pptx(file_bytes)
    assert parsed["format"] == "pptx"
    assert parsed["slide_count"] == 1
    assert "Executive Pitch Deck" in parsed["text"]
    assert "revolutionary engine" in parsed["text"]


def test_pdf_parser_fallback():
    # Test valid minimal PDF
    pdf_raw = (
        b"%PDF-1.4\n"
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R/Resources<<>>>>endobj\n"
        b"xref\n0 4\n0000000000 65535 f\n0000000009 00000 n\n0000000052 00000 n\n0000000101 00000 n\n"
        b"trailer<</Size 4/Root 1 0 R>>\nstartxref\n178\n%%EOF\n"
    )
    parsed = extract_text_from_pdf(pdf_raw)
    assert parsed["format"] == "pdf"
    assert parsed["page_count"] == 1


def test_upload_file_docx_endpoint():
    import docx
    doc = docx.Document()
    doc.add_heading("Global Q4 Hyper-Cloud Launch Announcement", level=1)
    doc.add_paragraph("BrandSphere is officially the fastest cloud platform in the world! Guaranteed 100% zero downtime.")
    buf = io.BytesIO()
    doc.save(buf)

    response = client.post(
        "/api/documents/upload-file",
        files={"file": ("launch_deck.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={
            "title": "Global Q4 Docx Launch",
            "author": "Enterprise Strategy",
            "content_type": "document"
        }
    )
    assert response.status_code == 201
    data = response.json()
    assert "document" in data
    assert data["document"]["title"] == "Global Q4 Docx Launch"
    assert data["issues_count"] > 0
    assert "metadata" in data
    assert data["metadata"]["format"] == "docx"


def test_upload_file_pptx_endpoint():
    import pptx
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Keynote Presentation"
    slide.placeholders[1].text = "Totally amazing platform that outshines every single competitor by a million miles."
    buf = io.BytesIO()
    prs.save(buf)

    response = client.post(
        "/api/documents/upload-file",
        files={"file": ("keynote.pptx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        data={
            "title": "Keynote Presentation Deck",
            "author": "Keynote Speaker",
            "content_type": "presentation"
        }
    )
    assert response.status_code == 201
    data = response.json()
    assert data["document"]["title"] == "Keynote Presentation Deck"
    assert data["issues_count"] > 0
    assert data["metadata"]["format"] == "pptx"

"""
BrandSphere AI - Document Parser Engine
Extracts plain text and structural metadata from PDF, DOCX, and PPTX binary documents
for pre-publication brand governance scanning.
"""

import io
from typing import Dict, Any, List


def extract_text_from_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from multi-page PDF documents using pypdf.
    Returns extracted text, page count, and page-by-page breakdown.
    """
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    page_count = len(reader.pages)
    pages_text: List[str] = []

    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        pages_text.append(page_text.strip())

    full_text = "\n\n".join(p for p in pages_text if p)
    return {
        "format": "pdf",
        "page_count": page_count,
        "text": full_text,
        "pages": pages_text,
        "paragraph_count": len([p for p in full_text.split("\n\n") if p.strip()])
    }


def extract_text_from_docx(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from Microsoft Word (.docx) documents using python-docx.
    Extracts body paragraphs, headings, and table cells.
    """
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs: List[str] = []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)

    # Also inspect tables for claims / metric data
    table_texts: List[str] = []
    for table in doc.tables:
        for row in table.rows:
            row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_vals:
                table_texts.append(" | ".join(row_vals))

    combined_text = "\n\n".join(paragraphs)
    if table_texts:
        combined_text += "\n\n[Document Tables]\n" + "\n".join(table_texts)

    return {
        "format": "docx",
        "page_count": max(1, len(paragraphs) // 5),  # Estimated page length based on paragraphs
        "paragraph_count": len(paragraphs),
        "text": combined_text,
        "table_count": len(doc.tables)
    }


def extract_text_from_pptx(file_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts text from Microsoft PowerPoint (.pptx) decks using python-pptx.
    Extracts slide titles, body text blocks, and presenter notes.
    """
    import pptx

    prs = pptx.Presentation(io.BytesIO(file_bytes))
    slide_count = len(prs.slides)
    slides_text: List[str] = []

    for i, slide in enumerate(prs.slides):
        slide_parts: List[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        # Check for slide notes if available
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(f"[Presenter Note: {notes}]")
        if slide_parts:
            slides_text.append(f"--- Slide {i + 1} ---\n" + "\n".join(slide_parts))

    full_text = "\n\n".join(slides_text)
    return {
        "format": "pptx",
        "page_count": slide_count,
        "slide_count": slide_count,
        "text": full_text,
        "paragraph_count": len(slides_text)
    }


def parse_document_file(file_name: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Dispatcher to extract text and metadata based on file extension.
    Supported: .pdf, .docx, .pptx, .txt, .md, .json
    """
    lower_name = file_name.lower()

    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif lower_name.endswith(".pptx"):
        return extract_text_from_pptx(file_bytes)
    elif lower_name.endswith((".txt", ".md", ".json")):
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")
        return {
            "format": lower_name.split(".")[-1],
            "page_count": 1,
            "paragraph_count": len([p for p in text.split("\n\n") if p.strip()]),
            "text": text
        }
    else:
        # Fallback to UTF-8 decoding
        try:
            text = file_bytes.decode("utf-8")
            fmt = "text"
        except Exception:
            text = file_bytes.decode("latin-1", errors="ignore")
            fmt = "binary"
        return {
            "format": fmt,
            "page_count": 1,
            "paragraph_count": 1,
            "text": text
        }
